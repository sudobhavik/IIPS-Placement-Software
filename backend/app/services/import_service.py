import logging
import re
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import openpyxl
from email_validator import EmailNotValidError, validate_email
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.enums import AcademicLevel, VerificationStatus
from app.models.import_job import ImportJob, ImportJobError
from app.models.student import Student, StudentAcademic
from app.models.user import User

logger = logging.getLogger(__name__)


ENROLLMENT_PATTERN = re.compile(r"^D[A-Z]\d{7}$")
ROLL_NO_PATTERN = re.compile(r"^(IC|IT|CS|EC|EE|ME|CE)-2K\d{2}-\d{2}$", re.IGNORECASE)

COURSE_NORMALIZATION = {
    "MCA": "MCA",
    "MTECH IT": "MTECH_IT",
    "MTECH_IT": "MTECH_IT",
    "MTECH": "MTECH_IT",
}

GENDER_NORMALIZATION = {
    "m": "male",
    "male": "male",
    "f": "female",
    "female": "female",
    "o": "other",
    "other": "other",
    "prefer not to say": "prefer_not_to_say",
    "prefer_not_to_say": "prefer_not_to_say",
}


class ImportErrorCode:
    INVALID_ENROLLMENT = "invalid_enrollment"
    INVALID_ROLL_NO = "invalid_roll_no"
    INVALID_COURSE = "invalid_course"
    INVALID_EMAIL = "invalid_email"
    INVALID_PERSONAL_NO = "invalid_personal_no"
    INVALID_GUARDIAN_NO = "invalid_guardian_no"
    INVALID_DOB = "invalid_dob"
    INVALID_AGE = "invalid_age"
    INVALID_SCORE = "invalid_score"
    INVALID_BACKLOGS = "invalid_backlogs"
    INVALID_GENDER = "invalid_gender"
    DUPLICATE_ENROLLMENT = "duplicate_enrollment"
    DUPLICATE_ROLL_NO = "duplicate_roll_no"
    DUPLICATE_EMAIL = "duplicate_email"
    COURSE_NOT_FOUND = "course_not_found"
    BATCH_NOT_FOUND = "batch_not_found"
    USER_EXISTS = "user_exists"
    STUDENT_EXISTS = "student_exists"
    SKIPPED_EXISTING = "skipped_existing"


class ImportWarningCode:
    SCALE_CONVERTED = "scale_converted"
    DOB_FROM_EXCEL_DATE = "dob_from_excel_date"
    INVALID_GUARDIAN_NO = "invalid_guardian_no"
    UG_CGPA_BLANK_MTECH = "ug_cgpa_blank_mtech"
    ROLL_NO_NONCONFORMING = "roll_no_nonconforming"
    SKIPPED_EXISTING = "skipped_existing"


class StudentImportService:
    def __init__(
        self,
        db: Session,
        job: ImportJob,
        current_user_id: int,
        batch_label: str | None = None,
        passing_year: int | None = None,
        current_semester: int = 7,
        dry_run: bool = False,
    ):
        self.db = db
        self.job = job
        self.current_user_id = current_user_id
        self.batch_label = batch_label
        self.passing_year = passing_year
        self.current_semester = current_semester
        self.dry_run = dry_run
        self.errors: list[ImportJobError] = []
        self.warnings = 0
        self.imported = 0
        self.placed_flag_ignored = 0
        self.skipped_existing = 0

        self.seen_enrollments: set[str] = set()
        self.seen_roll_nos: set[str] = set()
        self.seen_emails: set[str] = set()

        self.course_cache: dict[str, int] = {}
        self.batch_cache: dict[str, int] = {}

    def run(self, file_path: str) -> dict[str, Any]:
        wb = openpyxl.load_workbook(file_path, data_only=True)

        if "Complete Data" not in wb.sheetnames:
            raise ValueError("Sheet 'Complete Data' not found")

        ws = wb["Complete Data"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

        # Drop forbidden columns immediately
        forbidden_cols = {"username", "password"}
        header_map = {h: i for i, h in enumerate(headers) if h and h not in forbidden_cols}

        required_headers = [
            "email",
            "fullname",
            "enrollment_no",
            "gender",
            "dob",
            "personal_no",
            "guardian_no",
            "course",
            "10th_percent",
            "12th_percent",
            "ug_cgpa",
            "current_cgpa",
            "backlogs",
            "roll_no",
            "is_placed",
            "is_debarred",
        ]
        for h in required_headers:
            if h not in header_map:
                raise ValueError(f"Required header '{h}' not found")

        aggregated_data: dict[str, dict[str, Any]] = {}
        enrollment_row_num: dict[str, int] = {}

        # 1. Process Complete Data first
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            enrollment_idx = header_map.get("enrollment_no")
            if enrollment_idx is None or enrollment_idx >= len(row) or not row[enrollment_idx]:
                continue
            
            enrollment = str(row[enrollment_idx].value).strip().upper()
            enrollment = re.sub(r"\s+", "", enrollment)
            if not enrollment:
                continue

            row_dict = {}
            for h, idx in header_map.items():
                if idx < len(row) and row[idx]:
                    row_dict[h] = row[idx].value
            
            aggregated_data[enrollment] = row_dict
            enrollment_row_num[enrollment] = row_idx

        # 2. Process other sheets to enrich data (like linkedin_url from other sheets)
        for sheet_name in wb.sheetnames:
            if sheet_name == "Complete Data":
                continue
            
            other_ws = wb[sheet_name]
            other_headers = [cell.value for cell in next(other_ws.iter_rows(min_row=1, max_row=1))]
            other_header_map = {h: i for i, h in enumerate(other_headers) if h and h not in forbidden_cols}
            
            enrollment_key = next((k for k in ["enrollment_no", "Enrollment Number (Example: DX2200123)", "Enrollment Number"] if k in other_header_map), None)
            
            if not enrollment_key:
                continue
                
            other_enrollment_idx = other_header_map[enrollment_key]
            for row in other_ws.iter_rows(min_row=2, values_only=False):
                if other_enrollment_idx >= len(row) or not row[other_enrollment_idx]:
                    continue
                    
                enrollment = str(row[other_enrollment_idx].value).strip().upper()
                enrollment = re.sub(r"\s+", "", enrollment)
                
                if enrollment in aggregated_data:
                    for h, idx in other_header_map.items():
                        if h == enrollment_key:
                            continue
                        if idx < len(row) and row[idx] and row[idx].value is not None:
                            aggregated_data[enrollment][h] = row[idx].value

        # 3. Process the unified rows
        for enrollment, row_dict in aggregated_data.items():
            self._process_row(row_dict, enrollment_row_num[enrollment])

        self.job.total_rows = row_idx - 1
        self.job.success_rows = self.imported
        self.job.failed_rows = len(self.errors)
        self.job.status = "completed"
        self.job.finished_at = datetime.now(UTC)
        if not self.dry_run:
            self.db.commit()

        return {
            "total_rows": self.job.total_rows,
            "imported_rows": self.imported,
            "failed_rows": len(self.errors),
            "warnings": self.warnings,
            "placed_flag_ignored": self.placed_flag_ignored,
            "skipped_existing": self.skipped_existing,
            "errors": self.errors,
        }

    def _process_row(self, row_dict: dict[str, Any], row_num: int) -> None:
        def get_val(key: str) -> Any:
            return row_dict.get(key)

        raw_email = get_val("email")
        raw_fullname = get_val("fullname")
        raw_enrollment = get_val("enrollment_no")
        raw_gender = get_val("gender")
        raw_dob = get_val("dob")
        raw_personal = get_val("personal_no")
        raw_guardian = get_val("guardian_no")
        raw_course = get_val("course")
        raw_tenth = get_val("10th_percent")
        raw_twelfth = get_val("12th_percent")
        raw_diploma = get_val("diploma_percent")
        raw_ug = get_val("ug_cgpa")
        raw_pg = get_val("pg_cgpa")
        raw_current = get_val("current_cgpa")
        raw_backlogs = get_val("backlogs")
        raw_roll = get_val("roll_no")
        raw_is_placed = get_val("is_placed")
        raw_is_debarred = get_val("is_debarred")
        
        raw_father = get_val("father")
        raw_mother = get_val("mother")
        raw_caste = get_val("caste")
        raw_address = get_val("home_address")

        raw_linkedin = get_val("linkedin_url") or get_val("Your LinkedIn Link")
        raw_github = get_val("github_url") or get_val("Your GitHub Link")
        raw_portfolio = get_val("portfolio_url")
        raw_placement_opt_in = get_val("Are you interested in Campus Placements?")
        raw_10th_board = get_val("Board of Education - 10th standard")
        raw_12th_board = get_val("Board of Education - 12th standard")

        if raw_is_placed:
            self.placed_flag_ignored += 1

        enrollment = self._clean_enrollment(raw_enrollment, row_num)
        if not enrollment:
            return

        email = self._clean_email(raw_email, row_num)
        if not email:
            return

        roll_no = self._clean_roll_no(raw_roll, row_num) if raw_roll else None

        course_code = self._normalize_course(raw_course, row_num)
        if not course_code:
            return

        course_id = self._get_course_id(course_code, row_num)
        if not course_id:
            return

        batch_id = self._get_batch_id(row_num)
        if not batch_id:
            return

        gender = self._clean_gender(raw_gender, row_num) if raw_gender else None
        dob = self._parse_dob(raw_dob, row_num)
        if not dob:
            return

        personal_no = self._clean_phone(raw_personal, row_num, is_personal=True)
        if not personal_no:
            return

        guardian_no = self._clean_phone(raw_guardian, row_num, is_personal=False)

        tenth_pct = self._parse_score(raw_tenth, "10th_percent", row_num)
        twelfth_pct = self._parse_score(raw_twelfth, "12th_percent", row_num)
        diploma_pct = self._parse_score(raw_diploma, "diploma_percent", row_num)
        ug_cgpa_raw = self._parse_cgpa_raw(
            raw_ug,
            "ug_cgpa",
            row_num,
            allow_blank_for_mtech=(course_code == "MTECH_IT"),
        )
        pg_cgpa_raw = self._parse_cgpa_raw(
            raw_pg,
            "pg_cgpa",
            row_num,
            allow_blank_for_mtech=False,
        )
        current_percentage = self._parse_cgpa(raw_current, "current_cgpa", row_num)
        backlogs = self._parse_backlogs(raw_backlogs, row_num)
        is_debarred = bool(raw_is_debarred)

        placement_opt_in = True
        if raw_placement_opt_in is not None:
            opt_str = str(raw_placement_opt_in).strip().lower()
            if opt_str in ("no", "n", "false"):
                placement_opt_in = False

        board_10th = str(raw_10th_board).strip() if raw_10th_board and str(raw_10th_board).strip() else None
        board_12th = str(raw_12th_board).strip() if raw_12th_board and str(raw_12th_board).strip() else None

        if self.errors and self.errors[-1].row_number == row_num:
            return

        # Check for existing user/student
        existing_user = self.db.execute(
            select(User).where(User.email == email)
        ).scalar_one_or_none()
        existing_student = self.db.execute(
            select(Student).where(Student.enrollment_no == enrollment)
        ).scalar_one_or_none()

        if existing_user or existing_student:
            if not self.dry_run:
                self._handle_existing_student(
                    row_num, enrollment, email, existing_user, existing_student
                )
            else:
                self._add_warning(
                    row_num,
                    "enrollment_no",
                    ImportWarningCode.SKIPPED_EXISTING,
                    f"Would skip existing student {enrollment}",
                )
                self.skipped_existing += 1
            return

        if self.dry_run:
            # In dry run, just count as would-be imported
            self.imported += 1
            return

        # Create user
        user = User(
            email=email,
            full_name=self._clean_name(raw_fullname),
            phone=personal_no,
            password_hash=None,
            role="student",
            is_active=True,
            must_change_password=True,
        )
        self.db.add(user)
        self.db.flush()

        # Create student
        student = Student(
            user_id=user.id,
            enrollment_no=enrollment,
            roll_no=roll_no or f"UNKNOWN_{enrollment}",
            course_id=course_id,
            specialization_id=None,
            batch_id=batch_id,
            current_semester=self.current_semester,
            current_percentage=(
                Decimal(str(current_percentage)) if current_percentage is not None else None
            ),
            active_backlogs=backlogs,
            gap_years=0,
            gender=gender,
            date_of_birth=dob,
            father_name=self._clean_name(raw_father) if raw_father else None,
            mother_name=self._clean_name(raw_mother) if raw_mother else None,
            caste=str(raw_caste).strip() if raw_caste else None,
            city=self._extract_city_state(str(raw_address).strip() if raw_address else None)[0],
            state=self._extract_city_state(str(raw_address).strip() if raw_address else None)[1],
            linkedin_url=str(raw_linkedin).strip() if raw_linkedin else None,
            github_url=str(raw_github).strip() if raw_github else None,
            portfolio_url=str(raw_portfolio).strip() if raw_portfolio else None,
            guardian_phone=guardian_no,
            placement_opt_in=placement_opt_in,
            verification_status=VerificationStatus.PENDING,
            verified_by_id=None,
            verified_at=None,
            verification_remarks=None,
            profile_completeness=0,
            is_debarred=is_debarred,
        )
        self.db.add(student)
        self.db.flush()

        # Create academic records
        if tenth_pct is not None:
            self.db.add(
                StudentAcademic(
                    student_id=student.id,
                    level=AcademicLevel.TENTH,
                    institution=board_10th,
                    stream=None,
                    year_of_passing=None,
                    percentage=Decimal(str(tenth_pct)),
                    cgpa=None,
                )
            )
        if twelfth_pct is not None:
            self.db.add(
                StudentAcademic(
                    student_id=student.id,
                    level=AcademicLevel.TWELFTH,
                    institution=board_12th,
                    stream=None,
                    year_of_passing=None,
                    percentage=Decimal(str(twelfth_pct)),
                    cgpa=None,
                )
            )
        if diploma_pct is not None:
            self.db.add(
                StudentAcademic(
                    student_id=student.id,
                    level=AcademicLevel.DIPLOMA,
                    institution=None,
                    stream=None,
                    year_of_passing=None,
                    percentage=Decimal(str(diploma_pct)),
                    cgpa=None,
                )
            )
        if ug_cgpa_raw is not None:
            self.db.add(
                StudentAcademic(
                    student_id=student.id,
                    level=AcademicLevel.GRADUATION,
                    institution=None,
                    stream=None,
                    year_of_passing=None,
                    percentage=None,
                    cgpa=Decimal(str(ug_cgpa_raw)),
                )
            )
        if pg_cgpa_raw is not None:
            self.db.add(
                StudentAcademic(
                    student_id=student.id,
                    level=AcademicLevel.POST_GRADUATION,
                    institution=None,
                    stream=None,
                    year_of_passing=None,
                    percentage=None,
                    cgpa=Decimal(str(pg_cgpa_raw)),
                )
            )

        # Write import audit entry
        self._write_import_audit(user.id, student.id)

        self.db.commit()
        self.imported += 1

        logger.info("Imported student row %d: enrollment=%s", row_num, enrollment)

    def _handle_existing_student(
        self,
        row_num: int,
        enrollment: str,
        email: str,
        existing_user: User | None,
        existing_student: Student | None,
    ) -> None:
        """Handle idempotent re-run: skip existing, report differing fields."""
        differing_fields = []
        if existing_user:
            if existing_user.email != email:
                differing_fields.append("email")
        if existing_student:
            if existing_student.course_id != self.course_cache.get(enrollment):
                differing_fields.append("course")

        self.skipped_existing += 1
        msg = f"Skipped existing student {enrollment}"
        if differing_fields:
            msg += f"; differing fields: {', '.join(differing_fields)}"

        self._add_warning(
            row_num,
            "enrollment_no",
            ImportWarningCode.SKIPPED_EXISTING,
            msg,
        )

    def _write_import_audit(self, user_id: int, student_id: int) -> None:
        """Write an audit entry for the import with IDs only (no PII)."""
        from app.models.audit import AuditLog

        audit = AuditLog(
            actor_user_id=self.current_user_id,
            action="import",
            entity_type="student",
            entity_id=str(student_id),
            old_values=None,
            new_values={"user_id": user_id, "student_id": student_id},
            ip_address=None,
            request_id=None,
        )
        self.db.add(audit)

    def _clean_enrollment(self, value: Any, row_num: int) -> str | None:
        if value is None:
            self._add_error(
                row_num,
                "enrollment_no",
                ImportErrorCode.INVALID_ENROLLMENT,
                "Enrollment number is required",
            )
            return None
        cleaned = str(value).strip().upper()
        cleaned = re.sub(r"\s+", "", cleaned)
        if not ENROLLMENT_PATTERN.match(cleaned):
            self._add_error(
                row_num,
                "enrollment_no",
                ImportErrorCode.INVALID_ENROLLMENT,
                "Invalid enrollment format",
            )
            return None
        if cleaned in self.seen_enrollments:
            self._add_error(
                row_num,
                "enrollment_no",
                ImportErrorCode.DUPLICATE_ENROLLMENT,
                "Duplicate enrollment number in file",
            )
            return None
        self.seen_enrollments.add(cleaned)
        return cleaned

    def _clean_email(self, value: Any, row_num: int) -> str | None:
        if value is None:
            self._add_error(row_num, "email", ImportErrorCode.INVALID_EMAIL, "Email is required")
            return None
        cleaned = str(value).strip().lower()
        try:
            validate_email(cleaned, check_deliverability=False)
        except EmailNotValidError:
            self._add_error(row_num, "email", ImportErrorCode.INVALID_EMAIL, "Invalid email format")
            return None
        if cleaned in self.seen_emails:
            self._add_error(
                row_num,
                "email",
                ImportErrorCode.DUPLICATE_EMAIL,
                "Duplicate email in file",
            )
            return None
        self.seen_emails.add(cleaned)
        return cleaned

    def _clean_roll_no(self, value: Any, row_num: int) -> str | None:
        if value is None:
            return None
        cleaned = str(value).strip().upper()
        cleaned = re.sub(r"\s+", "-", cleaned)
        if not ROLL_NO_PATTERN.match(cleaned):
            self._add_warning(
                row_num,
                "roll_no",
                ImportWarningCode.ROLL_NO_NONCONFORMING,
                f"Roll number {cleaned} does not match expected pattern, importing as-is",
            )
            # Still add to seen to catch duplicates
        if cleaned in self.seen_roll_nos:
            self._add_error(
                row_num,
                "roll_no",
                ImportErrorCode.DUPLICATE_ROLL_NO,
                "Duplicate roll number in file",
            )
            return None
        self.seen_roll_nos.add(cleaned)
        return cleaned

    def _normalize_course(self, value: Any, row_num: int) -> str | None:
        if value is None:
            self._add_error(row_num, "course", ImportErrorCode.INVALID_COURSE, "Course is required")
            return None
        cleaned = str(value).strip().upper().replace(" ", "_")
        normalized = COURSE_NORMALIZATION.get(cleaned)
        if not normalized:
            self._add_error(
                row_num,
                "course",
                ImportErrorCode.INVALID_COURSE,
                f"Unknown course: {value}",
            )
            return None
        return normalized

    def _get_course_id(self, course_code: str, row_num: int) -> int | None:
        if course_code in self.course_cache:
            return self.course_cache[course_code]
        from app.models.master import Course

        course = self.db.execute(
            select(Course).where(Course.code == course_code, Course.is_active)
        ).scalar_one_or_none()
        if not course:
            self._add_error(
                row_num,
                "course",
                ImportErrorCode.COURSE_NOT_FOUND,
                f"Course {course_code} not found in master data",
            )
            return None
        self.course_cache[course_code] = course.id
        return course.id

    def _get_batch_id(self, row_num: int) -> int | None:
        cache_key = f"{self.batch_label}:{self.passing_year}"
        if cache_key in self.batch_cache:
            return self.batch_cache[cache_key]
        from app.models.master import Batch

        if self.batch_label and self.passing_year:
            batch = self.db.execute(
                select(Batch).where(
                    Batch.label == self.batch_label,
                    Batch.passing_year == self.passing_year,
                    Batch.is_active,
                )
            ).scalar_one_or_none()
        else:
            batch = self.db.execute(
    select(Batch).where(Batch.is_active).order_by(Batch.passing_year.desc()).limit(1)
).scalar_one_or_none()
        if not batch:
            self._add_error(
                row_num,
                "batch",
                ImportErrorCode.BATCH_NOT_FOUND,
                "No active batch found",
            )
            return None
        self.batch_cache[cache_key] = batch.id
        return batch.id

    def _clean_gender(self, value: Any, row_num: int) -> str | None:
        if value is None:
            return None
        cleaned = str(value).strip().lower()
        normalized = GENDER_NORMALIZATION.get(cleaned)
        if not normalized:
            self._add_warning(
                row_num,
                "gender",
                ImportWarningCode.INVALID_GENDER,
                f"Unknown gender value: {value}",
            )
            return None
        return normalized

    def _parse_dob(self, value: Any, row_num: int) -> date | None:
        if value is None or (isinstance(value, str) and value.strip() == ""):
            return None

        if isinstance(value, datetime):
            self._add_warning(
                row_num,
                "dob",
                ImportWarningCode.DOB_FROM_EXCEL_DATE,
                "DOB parsed from Excel datetime",
            )
            dob = value.date()
        elif isinstance(value, date):
            dob = value
        else:
            try:
                dob = datetime.strptime(str(value).strip(), "%d/%m/%Y").replace(tzinfo=UTC).date()
            except ValueError:
                self._add_error(
                    row_num,
                    "dob",
                    ImportErrorCode.INVALID_DOB,
                    "Invalid DOB format, expected dd/mm/yyyy",
                )
                return None

        today = datetime.now(UTC).date()
        age = today.year - dob.year - ((today.month, today.day) < (dob.month, dob.day))
        if age < 18 or age > 40:
            self._add_error(
                row_num,
                "dob",
                ImportErrorCode.INVALID_AGE,
                f"Age {age} outside valid range 18-40",
            )
            return None

        return dob

    def _clean_phone(self, value: Any, row_num: int, is_personal: bool) -> str | None:
        if value is None:
            if is_personal:
                self._add_error(
                    row_num,
                    "personal_no",
                    ImportErrorCode.INVALID_PERSONAL_NO,
                    "Personal number is required",
                )
                return None
            return None

        if isinstance(value, float):
            value = int(value)
        digits = re.sub(r"\D", "", str(value))
        if digits.startswith("91") and len(digits) == 12:
            digits = digits[2:]
        elif digits.startswith("0") and len(digits) == 11:
            digits = digits[1:]

        if len(digits) != 10 or digits[0] not in "6789":
            if is_personal:
                self._add_error(
                    row_num,
                    "personal_no",
                    ImportErrorCode.INVALID_PERSONAL_NO,
                    "Invalid personal number format",
                )
                return None
            else:
                self._add_warning(
                    row_num,
                    "guardian_no",
                    ImportWarningCode.INVALID_GUARDIAN_NO,
                    "Invalid guardian number format",
                )
                return None

        return digits

    def _parse_score(
        self,
        value: Any,
        field_name: str,
        row_num: int,
        allow_blank_for_mtech: bool = False,
    ) -> float | None:
        if value is None or (isinstance(value, str) and value.strip() == ""):
            if allow_blank_for_mtech:
                self._add_warning(
                    row_num,
                    field_name,
                    ImportWarningCode.UG_CGPA_BLANK_MTECH,
                    "UG CGPA blank for M.Tech student",
                )
            return None

        try:
            num = float(value)
        except (ValueError, TypeError):
            self._add_error(
                row_num,
                field_name,
                ImportErrorCode.INVALID_SCORE,
                f"Invalid {field_name} value",
            )
            return None

        if num <= 10:
            self._add_warning(
                row_num,
                field_name,
                ImportWarningCode.SCALE_CONVERTED,
                f"Value {num} treated as CGPA, multiplied by 10",
            )
            num *= 10

        if num < 0 or num > 100:
            self._add_error(
                row_num,
                field_name,
                ImportErrorCode.INVALID_SCORE,
                f"{field_name} {num} out of range 0-100",
            )
            return None

        return round(num, 2)

    def _parse_cgpa(
        self, value: Any, field_name: str, row_num: int, allow_blank_for_mtech: bool = False
    ) -> float | None:
        """Parse CGPA value. Returns percentage (0-100). If input is <=10, it's treated as CGPA and multiplied by 10."""
        if value is None or (isinstance(value, str) and value.strip() == ""):
            if allow_blank_for_mtech:
                self._add_warning(
                    row_num,
                    field_name,
                    ImportWarningCode.UG_CGPA_BLANK_MTECH,
                    "UG CGPA blank for M.Tech student",
                )
            return None

        try:
            num = float(value)
        except (ValueError, TypeError):
            self._add_error(
                row_num,
                field_name,
                ImportErrorCode.INVALID_SCORE,
                f"Invalid {field_name} value",
            )
            return None

        if num <= 10:
            num *= 10
        elif num > 100 or num < 0:
            self._add_error(
                row_num,
                field_name,
                ImportErrorCode.INVALID_SCORE,
                f"{field_name} {num} out of range 0-100",
            )
            return None

        return round(num, 2)

    def _parse_cgpa_raw(
        self, value: Any, field_name: str, row_num: int, allow_blank_for_mtech: bool = False
    ) -> float | None:
        """Parse CGPA value. Returns raw CGPA (0-10). If input > 10, it's treated as percentage and divided by 10."""
        if value is None or (isinstance(value, str) and value.strip() == ""):
            self._add_warning(
                row_num,
                field_name,
                "blank_score",
                f"{field_name} is blank",
            )
            return None

        try:
            num = float(value)
        except (ValueError, TypeError):
            self._add_error(
                row_num,
                field_name,
                ImportErrorCode.INVALID_SCORE,
                f"Invalid {field_name} value",
            )
            return None

        if num > 10 and num <= 100:
            num /= 10
        elif num > 100 or num < 0:
            self._add_error(
                row_num,
                field_name,
                ImportErrorCode.INVALID_SCORE,
                f"{field_name} {num} out of range 0-10",
            )
            return None

        return round(num, 2)

    def _parse_backlogs(self, value: Any, row_num: int) -> int:
        if value is None or (isinstance(value, str) and value.strip() == ""):
            return 0
        try:
            num = int(float(value))
        except (ValueError, TypeError):
            self._add_error(
                row_num,
                "backlogs",
                ImportErrorCode.INVALID_BACKLOGS,
                "Backlogs must be an integer",
            )
            return 0
        if num < 0:
            self._add_error(
                row_num,
                "backlogs",
                ImportErrorCode.INVALID_BACKLOGS,
                "Backlogs cannot be negative",
            )
            return 0
        return num

    def _clean_name(self, value: Any) -> str:
        if value is None:
            return ""
        cleaned = str(value).strip()
        cleaned = re.sub(r"\s+", " ", cleaned)
        return cleaned

    def _add_error(self, row_num: int, column: str | None, code: str, message: str) -> None:
        error = ImportJobError(
            job_id=self.job.id,
            row_number=row_num,
            column_name=column,
            message=f"[{code}] {message}",
            raw_row={"enrollment_no": None, "course": None},
        )
        self.db.add(error)
        self.db.flush()
        self.errors.append(error)
        logger.warning("Import error row %d col %s: %s", row_num, column, code)

    def _add_warning(self, row_num: int, column: str | None, code: str, message: str) -> None:
        self.warnings += 1
        logger.info("Import warning row %d col %s: %s", row_num, column, code)

    def _extract_city_state(self, address: str | None) -> tuple[str | None, str | None]:
        if not address:
            return None, None
        parts = [p.strip() for p in address.split(",") if p.strip()]
        if len(parts) >= 3:
            return parts[-2], parts[-1]
        elif len(parts) == 2:
            return parts[-1], None
        else:
            words = address.split()
            if words:
                return words[-1], None
        return None, None
